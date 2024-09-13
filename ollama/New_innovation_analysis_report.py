import ollama as ol
import json
import requests
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
import io
from bs4 import BeautifulSoup
import time
from datetime import datetime

def generate_patent_novelty_report(x):
    def remove_symbols_from_word(doc_bytes):
        buffer = io.BytesIO(doc_bytes)
        original_doc = Document(buffer)
        new_doc = Document()
        for paragraph in original_doc.paragraphs:
            clean_text = re.sub(r'[\*\#\-\[\]]', '', paragraph.text)
            new_doc.add_paragraph(clean_text)
        output_buffer = io.BytesIO()
        new_doc.save(output_buffer)
        output_buffer.seek(0)
        return output_buffer.getvalue()


    def search_patents(x1):
        post_url = 'https://www.chonghus.com/hxapi/qc/zytxt'
        headers = {
            'User-Agent': 'Mozilla/5.0',
            'Content-Type': 'application/json'
        }
        data = {
            'type': 'search_zy_txt',
            'Context': f"""{x1}""",
            'total': '10',
            'title': '摘要全文查重'
        }
        try:
            response = requests.post(post_url, headers=headers, json=data, timeout=10)
            time.sleep(1)
            response.raise_for_status()
        except requests.exceptions.RequestException as e:
            print(f"请求失败: {e}")
            exit()
        if response.status_code == 200:
            soup = BeautifulSoup(response.content, 'lxml')
            json_text = soup.p.text
            data = json.loads(json_text)
            unique_patents = {}
            for index, item in enumerate(data['msg']):
                unique_patents[str(index)] = item
            temp_dict = {}
            for key, value in unique_patents.items():
                pat_name = value['pat_name']
                if pat_name not in temp_dict:
                    temp_dict[pat_name] = value
            unique_patents = {str(i): v for i, v in enumerate(temp_dict.values())}
            for key in unique_patents:
                uuid = unique_patents[key]['uuid']
                url = f'https://www.chonghus.com/hxapi2/pat/detail?id={uuid}'
                headers = {
                    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36'
                }
                response = requests.get(url, headers=headers)
                time.sleep(1)
                if response.status_code == 200:
                    response.encoding = 'utf-8'
                    json_data = response.json()
                    msg_list = json_data.get('msg', {})
                    response_id = msg_list.get('id')
                    for key, patent in unique_patents.items():
                        if patent.get('uuid') == response_id:
                            unique_patents[key].update(msg_list)
                            break
            return unique_patents
        else:
            return (f"请求失败，状态码：{response.status_code}, 响应内容：{response.text}")

    def remove_symbols_from_word(doc_bytes):
        buffer = io.BytesIO(doc_bytes)
        original_doc = Document(buffer)
        new_doc = Document()
        for paragraph in original_doc.paragraphs:
            clean_text = re.sub(r'[\*\#\-\[\]]', '', paragraph.text)
            new_doc.add_paragraph(clean_text)
        output_buffer = io.BytesIO()
        new_doc.save(output_buffer)
        output_buffer.seek(0)
        processed_doc_bytes = output_buffer.getvalue()
        return processed_doc_bytes

    # 初始化文档
    doc = Document()
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal'].element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    doc.add_picture('./sheet/SRIPPM.png', width=Inches(2.0), height=Inches(0.5))
    content = doc.add_heading('', level=1).add_run('专利申请前新创性检索报告')
    content.font.name = u'宋体'
    content._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

    now = datetime.now()
    paragraph = doc.add_paragraph(
        f'本报告基于发明人提交的技术交底书，按照国内专利提案预审工作要求及其中的专利申请新颖性创造性检索标准，'
        f'预审人员在{now.year}年{now.month}月{now.day}日进行检索后分析完成。'
    )
    for run in paragraph.runs:
        run.font.name = u'宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

    content_1 = doc.add_heading('', level=2).add_run('一、提案基本信息')
    content_1.font.name = u'宋体'
    content_1._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    paragraph_1 = doc.add_paragraph(
        "提案名称：()\n"
        "提案单位：()\n"
        "提案类型：发明/实用新型\n"
        "技术联系人信息：\n"
        "姓名：()，手机：()，邮件：()\n"
        "预审人员信息：\n"
        "姓名：()，手机：()，邮件：()\n"
    )
    for run in paragraph_1.runs:
        run.font.name = u'宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

    content_2 = doc.add_heading('', level=2).add_run('二、预审分析意见')
    content_2.font.name = u'宋体'
    content_2._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    paragraph_2 = doc.add_paragraph(
        "本提案的方案属于专利法的保护客体，基于目前检索结果初步分析后，预审人员认为本提案具备新颖性及创造性，"
        "结合专利布局策略及行业专利分布情况，专利提案涉及技术方案具有一定的专利申请布局价值，故将其通过预审。具体分析如下："
    )
    for run in paragraph_2.runs:
        run.font.name = u'宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

    content_2_1 = doc.add_heading('', level=2).add_run('2.1、现有技术及本提案技术方案介绍')
    content_2_1.font.name = u'宋体'
    content_2_1._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

    # 系统提示
    default_prompt = (
        "你是一个专利专业人员，请阅读以上专利内容，详细回答我的任何问题，并且用中文回答我，我需要撰写一篇专利申请前新创性检索报告，"
        "请回答以下问题，确保内容详尽，清晰，并符合专利交底书的撰写要求，根据现有内容回答，且字数尽量长，专注于数据，不要说与问题无关的话，不要自己创造问题，"
        "不要太宽泛，具体到细节领域，直接回答问题内容，以下是我的问题：\n"
        "本提案技术方案介绍（首先介绍现有技术，然后介绍本提案解决了现有技术的什么问题）：\n"
        "现有技术一：\n"
        "现有技术二：\n"
        "本专利与现有技术一的区别：\n"
        "本专利与现有技术二的区别：\n"
        "申请策略建议：\n"
        "三、专利评分(每项评分条件按十分制评分。创造性，是否难以绕过，侵权判断是否容易。)：\n"
    )
    merged_prompt = re.split(r'\n', default_prompt)

    # 初始化聊天历史
    chat_history = [{"role": "system", "content": merged_prompt[0]}]

    # 搜索专利
    response_0 = search_patents(x)
    if isinstance(response_0, str):
        return f"Error: {response_0}"

    technology_1 = ('现有技术一：' + '\n\n' + '专利名：' + response_0['0']['pat_name'] + '\n\n' + '专利号：' + response_0['0']['open_no'] + '\n\n' +
                    '内容：' + response_0['0']['pat_qlyqs'])
    technology_2 = ('现有技术二：' + '\n\n' + '专利名：' + response_0['1']['pat_name'] + '\n\n' + '专利号：' + response_0['1']['open_no'] + '\n\n' +
                    '内容：' + response_0['1']['pat_qlyqs'])

    # 处理每个提示
    for i in range(1, len(merged_prompt) - 1):
        if "本专利与现有技术一的区别" in merged_prompt[i]:
            input_prompt = "本专利：" + x + '\n' + technology_1 + '\n' + merged_prompt[0] + '\n' + merged_prompt[i]
        elif "现有技术一" in merged_prompt[i] and "本专利" not in merged_prompt[i]:
            input_prompt = technology_1 + '\n' + merged_prompt[0] + '\n' + merged_prompt[i]
        elif "本专利与现有技术二的区别" in merged_prompt[i]:
            input_prompt = "本专利：" + x + '\n' + technology_2 + '\n' + merged_prompt[0] + '\n' + merged_prompt[i]
        elif "现有技术二" in merged_prompt[i] and "本专利" not in merged_prompt[i]:
            input_prompt = technology_2 + '\n' + merged_prompt[0] + '\n' + merged_prompt[i]
        else:
            input_prompt = "本专利：" + x + '\n' + merged_prompt[0] + '\n' + merged_prompt[i]

        user_message = {"role": "user", "content": input_prompt}
        chat_history.append(user_message)

        if "现有技术一" in merged_prompt[i] and "本专利" not in merged_prompt[i]:
            answer = technology_1
        elif "现有技术二" in merged_prompt[i] and "本专利" not in merged_prompt[i]:
            answer = technology_2
        elif "专利评分" in merged_prompt[i] and "本专利" not in merged_prompt[i]:
            response = ol.chat(model="llama2", messages=chat_history)
            answer = response["message"]["content"] + "\n" + "注：每项评分条件按十分制评分。专利评分在评审中比重为45%，其中，创造性25%，是否难以绕过5%，侵权判断是否容易15%。"
        else:
            response = ol.chat(model="qwen2:0.5b", messages=chat_history)
            answer = response["message"]["content"]

        run = doc.add_heading('', level=2).add_run(f'{merged_prompt[i]}')
        run.font.size = Pt(14)
        run.bold = True
        run.font.name = u'宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
        doc.add_paragraph(answer)

        ai_message = {"role": "assistant", "content": answer}
        chat_history.append(ai_message)

    content_4 = doc.add_heading('', level=2).add_run('四、预审结论')
    content_4.font.name = u'宋体'
    content_4._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    paragraph_4 = doc.add_paragraph(
        "本提案通过预审。\n"
    )
    for run in paragraph_4.runs:
        run.font.name = u'宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

    content_5 = doc.add_heading('', level=2).add_run('五、附件信息')
    content_5.font.name = u'宋体'
    content_5._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    paragraph_5 = doc.add_paragraph(
        f"对比文件1({response_0['0']['open_no']})" + '\n' + f"对比文件2({response_0['1']['open_no']})\n"
    )
    for run in paragraph_5.runs:
        run.font.name = u'宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    # 保存文档
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    doc_bytes = buffer.getvalue()
    doc_bytes = remove_symbols_from_word(doc_bytes)

    return doc_bytes

# 使用示例
# result = generate_patent_novelty_report("Your patent description here")
# with open("patent_novelty_report.docx", "wb") as f:
#     f.write(result)