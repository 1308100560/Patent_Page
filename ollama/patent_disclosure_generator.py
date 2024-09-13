import ollama as ol
import json
import requests
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
import io
from bs4 import BeautifulSoup
import time

def generate_patent_disclosure(x):
    def search_patents(x1):
        post_url = 'https://www.chonghus.com/hxapi/qc/zytxt'
        headers = {
            'User-Agent': 'Mozilla/5.0',
            'Content-Type': 'application/json'
        }
        data = {
            'type': 'search_zy_txt',
            'Context': f"""{x1}""",
            'total': '10',
            'title': '摘要全文查重'
        }
        try:
            response = requests.post(post_url, headers=headers, json=data, timeout=10)
            time.sleep(1)
            response.raise_for_status()
        except requests.exceptions.RequestException as e:
            print(f"请求失败: {e}")
            exit()
        if response.status_code == 200:
            soup = BeautifulSoup(response.content, 'lxml')
            json_text = soup.p.text
            data = json.loads(json_text)
            unique_patents = {}
            for index, item in enumerate(data['msg']):
                unique_patents[str(index)] = item
            temp_dict = {}
            for key, value in unique_patents.items():
                pat_name = value['pat_name']
                if pat_name not in temp_dict:
                    temp_dict[pat_name] = value
            unique_patents = {str(i): v for i, v in enumerate(temp_dict.values())}
            for key in unique_patents:
                uuid = unique_patents[key]['uuid']
                url = f'https://www.chonghus.com/hxapi2/pat/detail?id={uuid}'
                headers = {
                    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36'
                }
                response = requests.get(url, headers=headers)
                time.sleep(1)
                if response.status_code == 200:
                    response.encoding = 'utf-8'
                    json_data = response.json()
                    msg_list = json_data.get('msg', {})
                    response_id = msg_list.get('id')
                    for key, patent in unique_patents.items():
                        if patent.get('uuid') == response_id:
                            unique_patents[key].update(msg_list)
                            break
            return unique_patents
        else:
            return (f"请求失败，状态码：{response.status_code}, 响应内容：{response.text}")

    def remove_symbols_from_word(doc_bytes):
        buffer = io.BytesIO(doc_bytes)
        original_doc = Document(buffer)
        new_doc = Document()
        for paragraph in original_doc.paragraphs:
            clean_text = re.sub(r'[\*\#\-\[\]]', '', paragraph.text)
            new_doc.add_paragraph(clean_text)
        output_buffer = io.BytesIO()
        new_doc.save(output_buffer)
        output_buffer.seek(0)
        processed_doc_bytes = output_buffer.getvalue()
        return processed_doc_bytes

    # 初始化文档
    doc = Document()
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal'].element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    doc.add_picture('./SRIPPM.png', width=Inches(2.0), height=Inches(0.5))
    content = doc.add_heading('', level=1).add_run('专利交底书')
    content.font.name = u'宋体'
    content._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

    # 系统提示
    default_prompt = (
        "你是一个专利专业人员，请阅读以上专利内容，详细回答我的任何问题，并且用中文回答我，我需要撰写一篇专利交底书，"
        "请回答以下问题，确保内容详尽，清晰，并符合专利交底书的撰写要求，根据现有内容回答，且字数尽量长，专注于数据，不要说与问题无关的话，不要自己创造问题，"
        "不要太宽泛，具体到细节领域，直接回答问题内容，以下是我的问题：\n"
        "发明名称、技术领域、特点：\n"
        "背景技术：\n"
        "发明内容（解决的问题、解决所需技术方案、达到的有益效果）：\n"
        "附图说明(流程图，结构图)：\n"
        "实施方式：\n"
        "实施例（一个包含所有所知数据的流程）：\n"
        "对比例（现有技术和本技术的对比）：\n"
    )
    merged_prompt = re.split(r'\n', default_prompt)

    # 初始化聊天历史
    chat_history = [{"role": "system", "content": merged_prompt[0]}]

    # 搜索专利
    response_0 = search_patents(x)
    if isinstance(response_0, str):
        return f"Error: {response_0}"

    technology_1 = ('专利名：' + response_0['0']['pat_name'] + '\n\n' + '专利号：' + response_0['0']['open_no'] + '\n\n' +
                    '内容：' + response_0['0']['zhaiyao'])
    technology_2 = ('专利名：' + response_0['1']['pat_name'] + '\n\n' + '专利号：' + response_0['1']['open_no'] + '\n\n' +
                    '内容：' + response_0['1']['zhaiyao'])

    # 处理每个提示
    for i in range(1, len(merged_prompt) - 1):
        if "背景技术" in merged_prompt[i]:
            input_prompt = technology_1 + '\n' + technology_2 + '\n' + merged_prompt[0] + '\n' + merged_prompt[i]
        elif "对比例" in merged_prompt[i]:
            input_prompt = ("现有技术：" + technology_1 + '\n\n' + technology_2 + '\n\n' +
                            "本技术：" + x + '\n\n' + merged_prompt[0] + '\n\n' + merged_prompt[i])
        else:
            input_prompt = x + '\n\n' + merged_prompt[0] + '\n\n' + merged_prompt[i]

        user_message = {"role": "user", "content": input_prompt}
        chat_history.append(user_message)

        if "背景技术" in merged_prompt[i]:
            answer = technology_1 + '\n' + '\n' + technology_2
        elif "附图说明" in merged_prompt[i]:
            answer = "流程图，结构图"
        else:
            response = ol.chat(model="qwen2:0.5b", messages=chat_history)
            answer = response["message"]["content"]

        run = doc.add_heading('', level=2).add_run(f'{merged_prompt[i]}')
        run.font.size = Pt(14)
        run.bold = True
        run.font.name = u'宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
        doc.add_paragraph(answer)

        ai_message = {"role": "assistant", "content": answer}
        chat_history.append(ai_message)

    # 保存文档
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    doc_bytes = buffer.getvalue()
    doc_bytes = remove_symbols_from_word(doc_bytes)

    return doc_bytes

# 使用示例
# result = generate_patent_disclosure("Your input here")
# with open("patent_disclosure.docx", "wb") as f:
#     f.write(result)
